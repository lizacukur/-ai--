"""Выгрузка: Excel-таблица и HTML-панель с кнопками WhatsApp / Telegram."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .contacts import is_mobile, telegram_link, whatsapp_link
from .models import Lead
from .scoring import priority

OFFER_NAMES = {"site": "Сайт + SEO + Директ", "ai": "ИИ-помощник по сервису", "leadgen": "Лидогенерация"}


def channels(lead: Lead) -> tuple[str, str]:
    """Куда писать: (ссылка WhatsApp, ссылка Telegram). Явные мессенджеры важнее мобильного."""
    mobile = next((p for p in lead.phones if is_mobile(p)), "")
    wa = lead.whatsapp[0] if lead.whatsapp else mobile
    tg = lead.telegram[0] if lead.telegram else mobile
    return (whatsapp_link(wa, lead.message) if wa else "", telegram_link(tg) if tg else "")


def to_excel(leads: list[Lead], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"
    headers = ["Приоритет", "Компания", "Город", "Ниша", "ЛПР", "Должность", "Точность ЛПР", "ИНН",
               "Что предлагаем", "Нужен сайт (0-100)", "Нужен ИИ (0-100)", "WhatsApp", "Telegram",
               "Телефоны", "Email", "Сайт", "Рейтинг", "Отзывы", "Проблемы сайта", "Сообщение",
               "Источник", "Статус"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for lead in leads:
        wa, tg = channels(lead)
        ws.append([
            priority(lead), lead.name, lead.city, lead.niche, lead.lpr_name, lead.lpr_post,
            lead.site_facts.get("lpr_confidence", ""), lead.inn, OFFER_NAMES.get(lead.offer, lead.offer),
            lead.site_score, lead.ai_score, "Написать" if wa else "", "Открыть" if tg else "",
            ", ".join(lead.phones), ", ".join(lead.emails), lead.website, lead.rating, lead.reviews,
            "\n".join(lead.site_issues), lead.message, lead.source, "",
        ])
        row = ws.max_row
        for col, link in ((12, wa), (13, tg), (16, lead.website)):
            if link:
                ws.cell(row, col).hyperlink = link
                ws.cell(row, col).style = "Hyperlink"
        for col in (19, 20):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")

    widths = [10, 28, 16, 22, 28, 20, 18, 14, 22, 10, 10, 11, 11, 24, 26, 28, 8, 8, 40, 70, 9, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def to_html(leads: list[Lead], path: Path) -> None:
    rows = []
    for lead in leads:
        wa, tg = channels(lead)
        rows.append({
            "id": lead.key(), "name": lead.name, "city": lead.city, "niche": lead.niche,
            "lpr": " · ".join(x for x in (lead.lpr_name, lead.lpr_post) if x),
            "offer": OFFER_NAMES.get(lead.offer, lead.offer), "priority": priority(lead),
            "issues": lead.site_issues, "message": lead.message, "wa": wa, "tg": tg,
            "site": lead.website, "phones": lead.phones, "emails": lead.emails,
        })
    data = json.dumps(rows, ensure_ascii=False).replace("</", "<\\/")
    path.write_text(PANEL.replace("__DATA__", data), encoding="utf-8")


PANEL = """<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Лиды</title>
<style>
:root{--bg:#f6f7f9;--card:#fff;--text:#1d2129;--muted:#667085;--line:#e4e7ec;--accent:#2f5597;--wa:#1f9d55;--tg:#2481cc}
@media (prefers-color-scheme:dark){:root{--bg:#111418;--card:#1a1f26;--text:#e8eaed;--muted:#9aa4b2;--line:#2a313b;--accent:#7aa2e8}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:15px/1.45 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}
header{position:sticky;top:0;background:var(--bg);padding:16px;border-bottom:1px solid var(--line);z-index:1}
h1{margin:0 0 8px;font-size:20px}.bar{display:flex;flex-wrap:wrap;gap:8px}
input,select{padding:8px 10px;border:1px solid var(--line);border-radius:8px;background:var(--card);color:var(--text);font:inherit}
main{max-width:960px;margin:0 auto;padding:16px;display:grid;gap:12px}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:16px}
.card.done{opacity:.55}.top{display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap}
.name{font-weight:600;font-size:16px}.meta{color:var(--muted);font-size:13px}
.badge{display:inline-block;padding:2px 8px;border-radius:99px;background:var(--line);font-size:12px;margin-right:4px}
ul{margin:8px 0;padding-left:18px;color:var(--muted);font-size:13px}
textarea{width:100%;min-height:150px;padding:10px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--text);font:inherit;resize:vertical}
.actions{display:flex;flex-wrap:wrap;gap:8px;margin-top:8px}
button,a.btn{border:0;border-radius:8px;padding:9px 14px;font:inherit;font-weight:600;cursor:pointer;text-decoration:none;color:#fff;background:var(--accent)}
.wa{background:var(--wa)}.tg{background:var(--tg)}.ghost{background:transparent;color:var(--text);border:1px solid var(--line)}
.toast{position:fixed;bottom:16px;left:50%;transform:translateX(-50%);background:#000c;color:#fff;padding:8px 14px;border-radius:8px;display:none}
</style></head><body>
<header><h1>Лиды: <span id="count"></span></h1>
<div class="bar"><input id="q" placeholder="Поиск по названию, городу, нише">
<select id="st"><option value="">Все статусы</option><option value="new">Не писали</option><option value="sent">Написали</option><option value="reply">Ответили</option><option value="call">Созвон</option><option value="no">Отказ</option></select></div></header>
<main id="list"></main><div class="toast" id="toast"></div>
<script>
const DATA=__DATA__;
const ST={new:"Не писали",sent:"Написали",reply:"Ответили",call:"Созвон",no:"Отказ"};
const load=k=>{try{return JSON.parse(localStorage.getItem("lg_"+k)||"null")}catch(e){return null}};
const save=(k,v)=>{try{localStorage.setItem("lg_"+k,JSON.stringify(v))}catch(e){}};
const status=id=>(load("st")||{})[id]||"new";
function setStatus(id,v){const s=load("st")||{};s[id]=v;save("st",s);render()}
function toast(t){const e=document.getElementById("toast");e.textContent=t;e.style.display="block";setTimeout(()=>e.style.display="none",2200)}
async function copy(t){try{await navigator.clipboard.writeText(t);return true}catch(e){return false}}
const esc=s=>String(s??"").replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
function render(){
  const q=document.getElementById("q").value.toLowerCase(),f=document.getElementById("st").value;
  const items=DATA.filter(d=>(!q||(d.name+d.city+d.niche).toLowerCase().includes(q))&&(!f||status(d.id)===f));
  document.getElementById("count").textContent=items.length+" из "+DATA.length;
  document.getElementById("list").innerHTML=items.map((d,i)=>`<div class="card ${status(d.id)!=="new"?"done":""}">
   <div class="top"><div><div class="name">${esc(d.name)}</div>
   <div class="meta">${esc(d.city)} · ${esc(d.niche)}${d.lpr?" · <b>"+esc(d.lpr)+"</b>":""}</div></div>
   <div><span class="badge">${esc(d.offer)}</span><span class="badge">приоритет ${d.priority}</span></div></div>
   ${d.issues.length?"<ul>"+d.issues.map(x=>"<li>"+esc(x)+"</li>").join("")+"</ul>":""}
   <div class="meta">${d.site?`<a href="${esc(d.site)}" target="_blank">${esc(d.site)}</a> · `:""}${esc(d.phones.join(", "))} ${esc(d.emails.join(", "))}</div>
   <textarea data-id="${esc(d.id)}">${esc(d.message)}</textarea>
   <div class="actions">
    ${d.wa?`<button class="wa" data-act="wa" data-i="${i}">WhatsApp</button>`:""}
    ${d.tg?`<button class="tg" data-act="tg" data-i="${i}">Telegram</button>`:""}
    <button class="ghost" data-act="copy" data-i="${i}">Копировать текст</button>
    <select data-act="status" data-id="${esc(d.id)}">${Object.entries(ST).map(([k,v])=>`<option value="${k}" ${status(d.id)===k?"selected":""}>${v}</option>`).join("")}</select>
   </div></div>`).join("");
  document.querySelectorAll("button[data-act]").forEach(b=>b.onclick=async()=>{
    const d=items[+b.dataset.i],text=document.querySelector(`textarea[data-id="${CSS.escape(d.id)}"]`).value;
    if(b.dataset.act==="wa"){window.open(d.wa.split("?")[0]+"?text="+encodeURIComponent(text),"_blank");setStatus(d.id,"sent")}
    if(b.dataset.act==="tg"){const ok=await copy(text);toast(ok?"Текст скопирован — вставьте в чат":"Скопируйте текст вручную");window.open(d.tg,"_blank");setStatus(d.id,"sent")}
    if(b.dataset.act==="copy"){toast(await copy(text)?"Скопировано":"Не удалось скопировать")}
  });
  document.querySelectorAll("select[data-act=status]").forEach(s=>s.onchange=()=>setStatus(s.dataset.id,s.value));
}
document.getElementById("q").oninput=render;document.getElementById("st").onchange=render;render();
</script></body></html>"""
