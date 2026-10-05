"""Проверка без интернета: подменяем ответы 2ГИС, DaData и сайта."""
import json
from pathlib import Path

import yaml
from openpyxl import load_workbook

from leadgen import audit, contacts, discover, enrich, export, messages, scoring, sources
from leadgen.__main__ import share_lpr
from leadgen.models import Lead

CFG = yaml.safe_load((Path(__file__).parent.parent / "config.yaml").read_text(encoding="utf-8"))

OLD_SITE = """<html><head><title>Стоматология</title></head><body>
<p>Звоните: 8 (912) 345-67-89, почта info@smile-dent.ru</p>
<a href="https://wa.me/79123456789">WhatsApp</a> <a href="https://t.me/smiledent">TG</a>
<a href="https://t.me/+AbCdEf">группа</a>
<footer>© 2019 ООО «Смайл» ИНН 7701234567</footer></body></html>"""

GOOD_SITE = """<html><head><title>Стоматология «Улыбка» в Казани — лечение и имплантация</title>
<meta name="viewport" content="width=device-width"><meta name="description" content="Клиника">
<script src="https://mc.yandex.ru/metrika/tag.js"></script><script src="//w.yclients.com/widget.js"></script>
<script src="//code.jivo.ru/widget/x"></script></head><body><h1>Улыбка</h1>© 2026</body></html>"""


class FakeResp:
    def __init__(self, data, status=200):
        self._data, self.status_code = data, status
        self.text = json.dumps(data)

    def json(self):
        return self._data


def test_phone_normalization():
    assert contacts.normalize_phone("8 (912) 345-67-89") == "+79123456789"
    assert contacts.normalize_phone("+7 495 123 45 67") == "+74951234567"
    assert contacts.normalize_phone("123") == ""
    assert contacts.is_mobile("+79123456789") and not contacts.is_mobile("+74951234567")


def test_audit_old_site_finds_issues_and_contacts():
    lead = audit.analyze_html(Lead(name="Смайл", website="http://smile-dent.ru"), OLD_SITE,
                              final_url="http://smile-dent.ru", load_sec=5.2)
    text = " ".join(lead.site_issues)
    for expected in ("Нет HTTPS", "Медленно", "телефон", "© 2019", "онлайн-записи", "Метрики"):
        assert expected in text, expected
    assert lead.phones == ["+79123456789"]
    assert lead.emails == ["info@smile-dent.ru"]
    assert lead.whatsapp == ["+79123456789"]
    assert lead.telegram == ["@smiledent"]          # приглашение в группу не попало
    assert lead.inn == "7701234567"


def test_audit_good_site_is_clean():
    lead = audit.analyze_html(Lead(name="Улыбка", website="https://u.ru"), GOOD_SITE, final_url="https://u.ru")
    assert lead.site_issues == []
    assert lead.site_facts["booking"] and lead.site_facts["chat"] and lead.site_facts["metrika"]


def test_2gis_parsing(monkeypatch):
    item = {
        "name": "Смайл", "full_address_name": "Казань, ул. Баумана, 1", "org": {"branch_count": 3},
        "reviews": {"general_rating": 4.4, "general_review_count": 320},
        "contact_groups": [{"contacts": [
            {"type": "phone", "value": "+79123456789", "text": "+7 912 345-67-89"},
            {"type": "website", "value": "http://link.2gis.ru/abc", "text": "smile-dent.ru"},
            {"type": "telegram", "value": "https://t.me/smiledent", "text": "@smiledent"},
        ]}],
    }
    pages = [FakeResp({"meta": {"code": 200}, "result": {"items": [item]}}),
             FakeResp({"meta": {"code": 404}})]
    monkeypatch.setattr(sources.HTTP, "get", lambda *a, **k: pages.pop(0))
    [lead] = sources.search_2gis("key", "стоматология", "Казань", "dental", 50)
    assert lead.website == "http://smile-dent.ru"
    assert lead.phones == ["+79123456789"] and lead.telegram == ["@smiledent"]
    assert (lead.rating, lead.reviews, lead.branches) == (4.4, 320, 3)


def test_dadata_by_inn(monkeypatch):
    found = {"suggestions": [{"data": {
        "inn": "7701234567", "type": "LEGAL", "name": {"short_with_opf": "ООО «Смайл»"},
        "management": {"name": "Иванова Мария Петровна", "post": "ГЕНЕРАЛЬНЫЙ ДИРЕКТОР"}}}]}
    monkeypatch.setattr(enrich.requests, "post", lambda *a, **k: FakeResp(found))
    lead = enrich.find_lpr(Lead(name="Смайл", city="Казань", inn="7701234567"), "key")
    assert lead.lpr_name == "Иванова Мария Петровна"
    assert lead.lpr_post == "Генеральный директор"
    assert lead.site_facts["lpr_confidence"].startswith("точно")


def _party(name, okved, address, director):
    return {"value": f'ООО "{name}"', "data": {
        "inn": "7800000000", "type": "LEGAL", "okved": okved, "name": {"short": name, "short_with_opf": f'ООО "{name}"'},
        "address": {"value": address}, "management": {"name": director, "post": "ГЕНЕРАЛЬНЫЙ ДИРЕКТОР"}}}


def test_dadata_picks_clinic_by_okved_and_address(monkeypatch):
    """Одноимённых ООО «Династия» в городе много: берём то, что лечит и сидит по адресу из 2ГИС."""
    found = {"suggestions": [
        _party("ДИНАСТИЯ", "68.20", "г Санкт-Петербург, пр-кт Стачек, д 72", "Риелтор Иван Иванович"),
        _party("ДИНАСТИЯ", "86.10", "г Санкт-Петербург, ул Ленина, д 5 литера б", "Полякова Галина Юрьевна"),
        _party("ДИНАСТИЯ", "86.10", "г Москва, ул Ленина, д 5", "Московский Пётр Петрович"),
    ]}
    bodies = []
    monkeypatch.setattr(enrich.requests, "post", lambda url, json, **k: bodies.append(json) or FakeResp(found))
    lead = enrich.find_lpr(Lead(name="Династия, медицинский центр", city="Санкт-Петербург",
                                address="Санкт-Петербург, улица Ленина, 5 лит Б"), "key", ("86",))
    assert lead.lpr_name == "Полякова Галина Юрьевна"
    assert lead.site_facts["lpr_confidence"].startswith("вероятно")
    assert bodies[0]["query"] == "Династия" and bodies[0]["locations"] == [{"kladr_id": "78"}]


def test_dadata_rejects_generic_partial_match(monkeypatch):
    """«Клиника Пирогова» не должна привязаться к случайному ООО «Клиника» в другом месте города."""
    found = {"suggestions": [_party("КЛИНИКА", "86.23", "г Санкт-Петербург, ул Садовая, д 1", "Кто-то Другой")]}
    monkeypatch.setattr(enrich.requests, "post", lambda *a, **k: FakeResp(found))
    lead = enrich.find_lpr(Lead(name="Клиника Пирогова, медицинский центр", city="Санкт-Петербург",
                                address="Санкт-Петербург, Большой проспект В.О., 49"), "key", ("86",))
    assert lead.lpr_name == ""


def test_site_candidates_and_branch_sharing():
    assert discover.candidates("Так и ходи", "Санкт-Петербург")[:2] == ["takihodi.ru", "takihodi.com"]
    assert "dinastiya-spb.ru" in discover.candidates("Династия", "Санкт-Петербург")
    assert discover.candidates("ААА", "Санкт-Петербург") == []
    found = Lead(name="Так и ходи, салон", city="Санкт-Петербург", lpr_name="Руденко Елена Александровна",
                 site_facts={"lpr_confidence": "точно (ИНН с сайта)"})
    other = Lead(name="Так и ходи, салон для кудрявых", city="Санкт-Петербург")
    share_lpr([found, other])
    assert other.lpr_name == "Руденко Елена Александровна"
    assert "филиал" in other.site_facts["lpr_confidence"]


def test_scoring_and_message_and_export(tmp_path):
    no_site = Lead(name="Ромашка", city="Москва", niche="beauty_salon", phones=["+79001112233"])
    audit.audit_site(no_site)
    clinic = audit.analyze_html(
        Lead(name="Смайл", city="Казань", niche="dental", website="http://smile-dent.ru",
             reviews=320, rating=4.4, lpr_name="Иванова Мария Петровна", lpr_post="Директор"),
        GOOD_SITE.replace("yclients", "").replace("jivo", ""), final_url="https://smile-dent.ru")

    scoring.score(no_site, CFG["niches"]["beauty_salon"])
    scoring.score(clinic, CFG["niches"]["dental"])
    assert no_site.offer == "site" and no_site.site_score >= 50
    assert clinic.offer == "ai" and clinic.ai_score >= 70

    for lead in (no_site, clinic):
        lead.texts = messages.build(lead, CFG["niches"][lead.niche], CFG)
        lead.audience = messages.audience(lead)
        lead.message = lead.texts[f"{lead.audience}_messenger"]
    assert clinic.audience == "lpr" and no_site.audience == "admin"   # директора у салона нет
    assert clinic.message.startswith("Добрый день, Мария Петровна!")
    assert "помогаю стоматологиям" in clinic.message and clinic.message.endswith("больше всего времени?")
    assert clinic.texts["lpr_email"].endswith("Telegram: https://t.me/lizaai_consult")
    assert "руководителю салона" in no_site.message and "₽" not in no_site.message
    for lead in (no_site, clinic):
        for kind, text in lead.texts.items():
            assert messages.check_message(text, kind) == [], (kind, text)
    wa, tg = export.channels(clinic)
    assert wa == "" and tg == ""                   # контактов нет — кнопок нет
    wa, tg = export.channels(no_site)
    assert wa.startswith("https://wa.me/79001112233?text=") and tg == "https://t.me/+79001112233"

    export.to_excel([clinic, no_site], tmp_path / "l.xlsx")
    export.to_html([clinic, no_site], tmp_path / "l.html")
    ws = load_workbook(tmp_path / "l.xlsx").active
    assert ws.max_row == 3 and ws.cell(2, 2).value == "Смайл"
    assert "Ромашка" in (tmp_path / "l.html").read_text(encoding="utf-8")


def test_messages_follow_outreach_rules():
    """Все тексты по всем нишам проходят проверку; администратору имя директора в дательном падеже."""
    for niche, niche_cfg in CFG["niches"].items():
        for lpr in ("", "Магомедова Пасихат Батировна", "Канашкин Андрей Александрович"):
            texts = messages.build(Lead(name="X", niche=niche, lpr_name=lpr), niche_cfg, CFG)
            for kind, text in texts.items():
                assert messages.check_message(text, kind) == [], (kind, text)
    texts = messages.build(Lead(name="X", lpr_name="Магомедова Пасихат Батировна"), CFG["niches"]["beauty_clinic"], CFG)
    assert "передать это сообщение Пасихат Батировне" in texts["admin_messenger"]
    assert "передать это письмо Пасихат Батировне" in texts["admin_email"]
    assert "помогаю клиникам" in texts["lpr_messenger"] and "+79781335245" in texts["lpr_email"]


def test_audience_rule():
    owner = Lead(name="Slay", lpr_name="Ромашов Андрей Михайлович", lpr_post="Индивидуальный предприниматель", branches=5)
    small = Lead(name="Эверест", lpr_name="Чернов Михаил Юрьевич", lpr_post="Генеральный директор", branches=1)
    chain = Lead(name="Медси", lpr_name="Конорев Дмитрий Васильевич", lpr_post="Генеральный директор", branches=12)
    state = Lead(name="Алмазов", lpr_name="Шляхто Евгений Владимирович", legal_name='ФГБУ "НМИЦ"', branches=1)
    hired = Lead(name="Первая семейная", lpr_name="Чумаков Дмитрий Евгеньевич", lpr_note="Чумаков — наёмный управляющий")
    assert [messages.audience(x) for x in (owner, small, chain, state, hired, Lead(name="X"))] == \
        ["lpr", "lpr", "admin", "admin", "admin", "admin"]


def test_check_message_catches_broken_text():
    bad = "Привет, Анна! Давайте закажите — это стоит 20 000 ₽? Ок?"
    problems = messages.check_message(bad, "lpr_messenger")
    for part in ("тире", "один вопрос", "привет", "₽", "закажите"):
        assert any(part in p for p in problems), part


def test_known_inn_from_csv(tmp_path, monkeypatch):
    import leadgen.__main__ as main
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "inn.csv").write_text("brand,city,inn\nSkin Buro,Санкт-Петербург,7840080124\n", encoding="utf-8")
    monkeypatch.setattr(main, "ROOT", tmp_path)
    lead = Lead(name="Skin Buro, клиника эстетической медицины", city="Санкт-Петербург")
    main.apply_known_inn([lead])
    assert lead.inn == "7840080124"


def test_max_link_and_known_sites(tmp_path, monkeypatch):
    lead = audit.analyze_html(Lead(name="Смайл", website="https://s.ru"),
                              '<a href="https://max.ru/u/f9LHodD0cOI">MAX</a><a href="https://max.ru/">app</a>',
                              final_url="https://s.ru")
    assert lead.max == ["https://max.ru/u/f9LHodD0cOI"]
    assert export.max_link(lead) == "https://max.ru/u/f9LHodD0cOI"

    import leadgen.__main__ as main
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "sites.csv").write_text("brand,city,website\nДинастия,Санкт-Петербург,dinastiya-spb.ru\n",
                                                 encoding="utf-8")
    monkeypatch.setattr(main, "ROOT", tmp_path)
    lead = Lead(name="Династия, медицинский центр", city="Санкт-Петербург")
    main.apply_known_sites([lead])
    assert lead.website == "https://dinastiya-spb.ru"


def test_lpr_profile_and_director_email(tmp_path, monkeypatch):
    import leadgen.__main__ as main
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "lpr_profiles.csv").write_text(
        "brand,city,profile,note\nКэнворк,Санкт-Петербург,https://kanwork.ru/kanashkin,Страница владельца\n",
        encoding="utf-8")
    monkeypatch.setattr(main, "ROOT", tmp_path)
    lead = Lead(name="Кэнворк, автосервис", city="Санкт-Петербург", lpr_name="Канашкин Андрей Александрович",
                emails=["mail@kanwork.ru"], message="Здравствуйте!")
    main.apply_lpr_profiles([lead])
    assert lead.lpr_profile == "https://kanwork.ru/kanashkin" and lead.lpr_note == "Страница владельца"
    assert export.email_subject(lead) == "Руководителю: Канашкин Андрей Александрович"
    export.to_excel([lead], tmp_path / "l.xlsx")
    export.to_html([lead], tmp_path / "l.html")
    ws = load_workbook(tmp_path / "l.xlsx").active
    assert ws.cell(1, 24).value == "Профиль ЛПР" and ws.cell(2, 25).value == "Страница владельца"
    html = (tmp_path / "l.html").read_text(encoding="utf-8")
    assert "Отправить на почту" in html and "kanwork.ru/kanashkin" in html
